import sys
import openpyxl


def create_pool_xlsx(output_path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "China RFBS"

    headers = [
        "Priority Group", "Delivery Method", "Info", "Name", "Notes",
        "Days", "Rate", "Battery", "Liquid", "Size",
        "Min Weight", "Max Weight", "Notes2", "Price Limit",
        "Col15", "Col16", "Volume Mode", "Volume Coeff",
    ]
    for col, val in enumerate(headers, start=1):
        ws.cell(5, col).value = val

    entries = [
        ("A", "land", "", "E2E_Pool_Test_1", "", "7-14", "10+2/100g", "No", "No", "", 0, 2000, "", "", "", "", "physical weight", "5000"),
        ("B", "air", "", "E2E_Pool_Test_2", "", "3-7", "15+3/100g", "No", "No", "", 0, 1000, "", "", "", "", "physical weight", "5000"),
        ("C", "sea", "", "E2E_Pool_Test_3", "", "20-30", "5+1/100g", "No", "No", "", 0, 5000, "", "", "", "", "physical weight", "5000"),
    ]

    for i, entry in enumerate(entries):
        row = 6 + i
        for col, val in enumerate(entry, start=1):
            ws.cell(row, col).value = val

    wb.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(1)
    create_pool_xlsx(sys.argv[1])
