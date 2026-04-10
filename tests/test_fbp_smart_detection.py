from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from xlsx_import.parsers import fbp


def _set_cell(sheet: Worksheet, row: int, column: int, value: Any) -> None:
    cell = cast(Any, sheet.cell(row, column))
    cell.value = value


def _save_fbp_workbook(tmp_path: Path, rows: list[dict[str, object]]) -> str:
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "China FBP"

    for offset, row in enumerate(rows, start=8):
        _set_cell(sheet, offset, 1, row.get("index", offset - 7))
        _set_cell(sheet, offset, 2, row.get("company", "CEL"))
        _set_cell(sheet, offset, 3, row.get("service", ""))
        _set_cell(sheet, offset, 4, row.get("type", "Standard"))
        _set_cell(sheet, offset, 5, row.get("destination", "Russia"))
        _set_cell(sheet, offset, 6, row.get("priority", "D"))
        _set_cell(sheet, offset, 7, row["name"])
        _set_cell(sheet, offset, 8, row.get("days", "5-7"))
        _set_cell(sheet, offset, 9, row.get("rate", "¥8 + ¥4.8/100g"))
        _set_cell(sheet, offset, 10, row.get("battery", "No"))
        _set_cell(sheet, offset, 11, row.get("liquid", "No"))
        _set_cell(sheet, offset, 12, row.get("size", "sum of sides 90, length 50"))
        _set_cell(sheet, offset, 13, row.get("min_weight", "0"))
        _set_cell(sheet, offset, 14, row.get("max_weight", "30000"))
        _set_cell(sheet, offset, 15, row.get("price_limit", "10-100"))

    filepath = tmp_path / "fbp-smart-detection.xlsx"
    workbook.save(filepath)
    return str(filepath)


def _parse_single_row(tmp_path: Path, service: str, name: str) -> dict:
    filepath = _save_fbp_workbook(
        tmp_path,
        [{"service": service, "name": name}],
    )
    result = fbp.parse(filepath)
    assert result["errors"] == []
    assert result["skipped"] == 0
    assert len(result["rows"]) == 1
    return result["rows"][0]


def test_fbp_marks_smart_when_name_contains_smart_but_service_does_not(tmp_path: Path):
    row = _parse_single_row(tmp_path, "CEL FBP Standard", "Special Smart Lane")

    assert row["is_smart"] == 1


def test_fbp_marks_smart_when_service_contains_smart_but_name_does_not(tmp_path: Path):
    row = _parse_single_row(tmp_path, "CEL FBP Smart", "CEL FBP Standard Extra Small")

    assert row["is_smart"] == 1


def test_fbp_marks_smart_when_service_and_name_both_contain_smart(tmp_path: Path):
    row = _parse_single_row(tmp_path, "CEL FBP Smart", "CEL Smart Premium")

    assert row["is_smart"] == 1


def test_fbp_marks_non_smart_when_service_and_name_both_lack_smart(tmp_path: Path):
    row = _parse_single_row(
        tmp_path, "CEL FBP Standard", "CEL FBP Standard Extra Small"
    )

    assert row["is_smart"] == 0


def test_fbp_smart_detection_is_case_insensitive_in_service_and_name(tmp_path: Path):
    filepath = _save_fbp_workbook(
        tmp_path,
        [
            {"service": "CEL FBP SMART", "name": "Service Variant"},
            {"service": "CEL FBP Standard", "name": "Name Smart Variant"},
            {"service": "CEL FBP smart", "name": "Lowercase Variant"},
        ],
    )

    result = fbp.parse(filepath)

    assert [row["is_smart"] for row in result["rows"]] == [1, 1, 1]


def test_fbp_carries_smart_service_value_across_merged_style_rows(tmp_path: Path):
    filepath = _save_fbp_workbook(
        tmp_path,
        [
            {
                "company": "CEL",
                "service": "CEL FBP Smart",
                "type": "Standard",
                "name": "First Service Row",
            },
            {
                "company": "",
                "service": "",
                "type": "",
                "name": "Second Carried Row",
            },
        ],
    )

    result = fbp.parse(filepath)

    assert len(result["rows"]) == 2
    assert result["rows"][0]["is_smart"] == 1
    assert result["rows"][1]["is_smart"] == 1


def test_fbp_reads_smart_from_merged_cells_ranges(tmp_path: Path):
    filepath = _save_fbp_workbook(
        tmp_path,
        [
            {
                "company": "CEL",
                "service": "Smart",
                "type": "Standard",
                "name": "Merged Smart First",
            },
            {
                "company": "",
                "service": "",
                "type": "",
                "name": "Merged Smart Second",
            },
        ],
    )
    saved = load_workbook(filepath)
    sheet = cast(Worksheet, saved.active)
    sheet.merge_cells("C8:C9")
    saved.save(filepath)

    result = fbp.parse(filepath)

    assert [row["is_smart"] for row in result["rows"]] == [1, 1]
