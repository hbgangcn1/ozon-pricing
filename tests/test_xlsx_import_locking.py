from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

import xlsx_import


def _create_workbook(path: Path, title: str = "CHINA FBP") -> str:
    workbook = Workbook()
    workbook.active.title = title
    workbook.save(path)
    return str(path)


def test_parse_file_opens_workbook_once_and_passes_same_workbook(
    tmp_path: Path, monkeypatch
):
    filepath = _create_workbook(tmp_path / "FBP_List_of_services_test.xlsx")
    real_loader = openpyxl.load_workbook
    seen: dict[str, object] = {}
    calls: list[str] = []

    def fake_loader(path: str, data_only: bool = True):
        calls.append(path)
        return real_loader(path, data_only=data_only)

    def fake_detect_format(source):
        seen["detector"] = source
        return "fbp"

    def fake_parser(source, filepath=None, source_file=None):
        seen["parser"] = source
        seen["filepath"] = filepath
        seen["source_file"] = source_file
        return {"rows": [], "skipped": 0, "errors": []}

    monkeypatch.setattr(xlsx_import.openpyxl, "load_workbook", fake_loader)
    monkeypatch.setattr(xlsx_import, "detect_format", fake_detect_format)
    monkeypatch.setitem(xlsx_import.PARSERS, "fbp", fake_parser)

    result = xlsx_import.parse_file(filepath, source_file="source.xlsx")

    assert result["format"] == "fbp"
    assert calls == [filepath]
    assert seen["detector"] is seen["parser"]
    assert seen["filepath"] == filepath
    assert seen["source_file"] == "source.xlsx"


def test_parse_file_retries_with_temp_copy_after_permission_error(
    tmp_path: Path, monkeypatch
):
    filepath = _create_workbook(tmp_path / "FBP_List_of_services_locked.xlsx")
    real_loader = openpyxl.load_workbook
    copied_from: list[tuple[str, str]] = []
    sleeps: list[float] = []
    calls: list[str] = []

    def fake_loader(path: str, data_only: bool = True):
        calls.append(path)
        if path == filepath:
            raise PermissionError("[WinError 32] file locked")
        return real_loader(path, data_only=data_only)

    def fake_copy2(src: str, dst: str):
        copied_from.append((src, dst))
        Path(dst).write_bytes(Path(src).read_bytes())
        return dst

    monkeypatch.setattr(xlsx_import.openpyxl, "load_workbook", fake_loader)
    monkeypatch.setattr(xlsx_import.shutil, "copy2", fake_copy2)
    monkeypatch.setattr(xlsx_import.time, "sleep", sleeps.append)

    result = xlsx_import.parse_file(filepath)

    assert result["format"] == "fbp"
    assert calls[0] == filepath
    assert len(copied_from) == 1
    assert copied_from[0][0] == filepath
    assert calls[-1] == copied_from[0][1]
    assert sleeps == [0.1]
    assert not Path(copied_from[0][1]).exists()


def test_parse_file_raises_last_permission_error_after_retry_budget(
    tmp_path: Path, monkeypatch
):
    filepath = _create_workbook(tmp_path / "FBP_List_of_services_fail.xlsx")
    copied: list[tuple[str, str]] = []
    sleeps: list[float] = []
    calls: list[str] = []

    def fake_loader(path: str, data_only: bool = True):
        calls.append(path)
        raise PermissionError(f"[WinError 32] {path}")

    def fake_copy2(src: str, dst: str):
        copied.append((src, dst))
        Path(dst).write_bytes(Path(src).read_bytes())
        return dst

    monkeypatch.setattr(xlsx_import.openpyxl, "load_workbook", fake_loader)
    monkeypatch.setattr(xlsx_import.shutil, "copy2", fake_copy2)
    monkeypatch.setattr(xlsx_import.time, "sleep", sleeps.append)

    with pytest.raises(PermissionError, match="WinError 32"):
        xlsx_import.parse_file(filepath)

    assert len(copied) == 1
    assert calls == [filepath, copied[0][1], copied[0][1], copied[0][1]]
    assert sleeps == [0.1, 0.2, 0.4]
    assert all(not Path(path).exists() for _, path in copied)
